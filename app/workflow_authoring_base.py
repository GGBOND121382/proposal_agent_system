from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from .context_base import (
    REPORT_CONTENT_CRITIC_PROMPT,
    REPORT_OUTLINE_PROMPT,
    REPORT_SECTION_WRITE_PROMPT,
)
from .executor import PromptExecutionError
from .llm import MODEL_RESPONSE_PROTOCOL_VERSION
from .retry_policy import ProviderRetriesExhausted
from .runtime_failures import FailureCategory, classify_runtime_failure
from .util import new_id, sha256_json, sha256_text, utc_now
from .workflow_input import CURRENT_PROPOSAL_INPUT, WorkflowInputRequired, material_input_questions
from .workflow_status import WorkflowStatus

THREE_SECTION_PROFILE_ORDER = (
    "BACKGROUND_AND_SIGNIFICANCE",
    "RESEARCH_CONTENT",
    "TECHNICAL_ROUTE",
)


class WorkflowAuthoringMixin:
    SECTION_PHASES = {
        "BLUEPRINT": ("P-WRITE-BLUEPRINT", "BLUEPRINT_CRITIC"),
        "BLUEPRINT_CRITIC": ("P-WRITE-BLUEPRINT-CRITIC", "CONTENT"),
        "CONTENT": ("P-WRITE-CONTENT", "CONTENT_CRITIC"),
        "CONTENT_CRITIC": ("P-WRITE-CRITIC", "POLISH"),
        "POLISH": ("P-EXPRESSION-POLISH", "EXPRESSION_CRITIC"),
        "EXPRESSION_CRITIC": ("P-EXPRESSION-CRITIC", "DONE"),
    }
    SECTION_REPAIR_CRITICS = {
        "P-WRITE-BLUEPRINT-CRITIC",
        "P-WRITE-CRITIC",
        "P-EXPRESSION-CRITIC",
    }
    SECTION_CRITIC_PRODUCERS = {
        "P-WRITE-BLUEPRINT-CRITIC": ("P-WRITE-BLUEPRINT", "BLUEPRINT"),
        "P-WRITE-CRITIC": ("P-WRITE-CONTENT", "CONTENT"),
        "P-EXPRESSION-CRITIC": ("P-EXPRESSION-POLISH", "POLISH"),
    }
    SECTION_PRODUCER_PHASES = {
        "P-WRITE-BLUEPRINT": "BLUEPRINT",
        "P-WRITE-CONTENT": "CONTENT",
        "P-EXPRESSION-POLISH": "POLISH",
    }

    @staticmethod
    def _append_section_run(progress: dict[str, Any], result: dict[str, Any], *, prompt_id: str | None = None, role: str | None = None) -> None:
        run_id = str(result.get("run_id") or "")
        if run_id and any(str(item.get("run_id") or "") == run_id for item in progress["runs"]):
            return
        record = {
            "prompt_id": prompt_id or result.get("prompt_id"),
            "run_id": run_id,
            "status": result.get("status"),
        }
        if role:
            record["role"] = role
        progress["runs"].append(record)

    def _apply_section_decision(
        self,
        wf: dict[str, Any],
        state: dict[str, Any],
        prompt_id: str,
        result: dict[str, Any],
    ) -> dict[str, Any]:
        """Derive section control state through the workflow decision authority.

        Prompt runs and their model outputs are immutable evidence.  The section
        chain must nevertheless branch on the same critic/guard arbitration as
        ordinary workflow steps; otherwise a model-level ``REVISE`` containing
        no owned blocking finding can consume repair budget and block the
        workflow.  Keep the raw result untouched and return a derived control
        view carrying the effective status and actionable findings.

        The callable check preserves the standalone mixin contract used by
        lightweight authoring harnesses.  The production WorkflowEngine always
        supplies ``_record_decision``.
        """
        recorder = getattr(self, "_record_decision", None)
        derived = dict(result)
        if callable(recorder):
            decision, effective_status, effective_output = recorder(
                wf,
                state,
                prompt_id,
                result,
            )
            derived["model_status"] = str(
                result.get("status")
                or (result.get("output") or {}).get("status")
                or "ERROR"
            )
            derived["status"] = effective_status
            derived["output"] = effective_output
            if decision is not None:
                derived["decision"] = decision
        self._sync_section_revision_feedback(state, prompt_id, derived)
        return derived

    def _sync_section_revision_feedback(
        self,
        state: dict[str, Any],
        prompt_id: str,
        result: dict[str, Any],
    ) -> None:
        """Keep producer feedback aligned with the latest effective decision.

        Revision feedback is control state, not an append-only audit stream.  A
        new REVISE/BLOCK decision must replace the previous candidate's feedback
        even when the regeneration budget is already exhausted.  Conversely, a
        producer PASS proves that the feedback supplied to that producer has
        been consumed and must not leak into a later authoring phase.  Raw runs
        and decision artifacts remain the immutable history.
        """
        section_id = str(state.get("active_section_id") or "")
        if not section_id:
            return
        status = str(
            result.get("status")
            or (result.get("output") or {}).get("status")
            or ""
        ).upper()
        feedback_by_section = state.setdefault("section_revision_findings", {})
        if status in {"REVISE", "BLOCK"}:
            output = result.get("output") or {}
            feedback_by_section[section_id] = [
                dict(item)
                for item in output.get("findings") or []
                if isinstance(item, dict)
            ]
        elif status == "PASS" and prompt_id in self.SECTION_PRODUCER_PHASES:
            feedback_by_section.pop(section_id, None)

    def _block_section_chain(
        self,
        wf: dict[str, Any],
        state: dict[str, Any],
        section: dict[str, Any],
        message: str,
        *,
        configuration_error: Exception | str | None = None,
    ) -> dict[str, Any]:
        repair_failure = state.get("last_targeted_repair_failure")
        if isinstance(repair_failure, dict):
            message = self._targeted_repair_failure_message(
                state,
                prompt_id=str(
                    repair_failure.get("critic_prompt")
                    or "P-TARGETED-REPAIR"
                ),
                fallback=message,
            )
            if (
                configuration_error is None
                and str(repair_failure.get("category") or "")
                == FailureCategory.CONFIGURATION.value
            ):
                configuration_error = str(
                    repair_failure.get("error") or message
                )
        section_id = str(section.get("section_id") or "")
        progress = state.setdefault("section_progress", {}).setdefault(section_id, {})
        report = (
            self._runtime_configuration_report(
                configuration_error,
                scope=f"SECTION_CHAIN_RUNTIME:{section_id or 'UNKNOWN'}",
            )
            if configuration_error is not None
            else None
        )
        if report is not None:
            progress["status"] = "WAITING_CONFIGURATION"
            progress["last_error"] = message
            return self._pause_for_configuration(
                wf,
                state,
                report,
                source=f"SECTION_CHAIN_RUNTIME:{section_id or 'UNKNOWN'}",
            )
        if isinstance(repair_failure, dict):
            blocked_status = self._targeted_repair_block_status(state)
        elif configuration_error is not None:
            blocked_status = WorkflowStatus.BLOCKED_TECHNICAL.value
        else:
            blocked_status = WorkflowStatus.BLOCKED_CONTENT.value
        progress["status"] = blocked_status
        progress["last_error"] = message
        state["last_error"] = f"{section.get('title')}: {message}"
        self._update(wf, status=blocked_status, state=state)
        return self.get(wf["id"])

    @staticmethod
    def _targeted_repair_block_status(state: dict[str, Any]) -> str:
        """Preserve the actual targeted-repair failure category.

        A failed repair used to be flattened into the historical generic
        ``BLOCKED`` state.  That hid whether the repair provider, output
        contract, runtime, or semantic content was responsible and prevented
        the contract-migration path from selecting the saved repair output.
        """

        failure = state.get("last_targeted_repair_failure")
        if not isinstance(failure, dict):
            return WorkflowStatus.BLOCKED_TECHNICAL.value
        category = str(failure.get("category") or "")
        mapping = {
            FailureCategory.OUTPUT_CONTRACT.value: WorkflowStatus.BLOCKED_CONTRACT.value,
            FailureCategory.PROVIDER_TRANSIENT.value: WorkflowStatus.BLOCKED_PROVIDER.value,
            FailureCategory.TECHNICAL.value: WorkflowStatus.BLOCKED_TECHNICAL.value,
            FailureCategory.SEMANTIC_REVISE.value: WorkflowStatus.BLOCKED_CONTENT.value,
            # If dependency preflight can identify the missing configuration,
            # _block_section_chain converts it to WAITING_CONFIGURATION before
            # this fallback.  Without such a report, keep the workflow paused
            # instead of auto-resuming an unqualified waiting state.
            FailureCategory.CONFIGURATION.value: WorkflowStatus.BLOCKED_TECHNICAL.value,
            "SEMANTIC_REPAIR_REJECTED": WorkflowStatus.BLOCKED_CONTENT.value,
        }
        return mapping.get(category, WorkflowStatus.BLOCKED_TECHNICAL.value)

    _section_block_status = _targeted_repair_block_status

    def _block_section_prompt_failure(
        self,
        wf: dict[str, Any],
        state: dict[str, Any],
        section: dict[str, Any],
        exc: PromptExecutionError,
    ) -> dict[str, Any]:
        """Block or pause a section using the shared runtime failure taxonomy.

        Legacy section checkpoints can surface schema failures as a plain
        ``PromptExecutionError(validation_errors=...)`` instead of a typed
        provider exception.  Do not flatten those contract failures into the
        historical technical-error fallback.
        """
        classification = classify_runtime_failure(exc)
        if classification.category is FailureCategory.CONFIGURATION:
            return self._block_section_chain(
                wf,
                state,
                section,
                str(exc),
                configuration_error=exc,
            )

        section_id = str(section.get("section_id") or "")
        progress = state.setdefault("section_progress", {}).setdefault(
            section_id, {}
        )
        status = classification.workflow_status
        progress["status"] = status
        progress["last_error"] = str(exc)
        state["last_error"] = f"{section.get('title')}: {exc}"
        self._update(wf, status=status, state=state)
        return self.get(wf["id"])

    async def _execute_section_prompt(
        self,
        wf: dict[str, Any],
        state: dict[str, Any],
        section: dict[str, Any],
        progress: dict[str, Any],
        prompt_id: str,
        *,
        role: str | None = None,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        envelope = self.context_builder.build(
            prompt_id,
            wf["project_id"],
            workflow_id=wf["id"],
            workflow_state=state,
        )
        candidate_round_key = (
            f"section:{state.get('active_section_id') or ''}:{prompt_id}"
        )
        candidate_round = int(
            (state.get("acceptance_candidate_rounds") or {}).get(
                candidate_round_key,
                0,
            )
        )
        requested_call_key = None
        if candidate_round:
            requested_call_key = "call-acceptance-" + sha256_json(
                {
                    "workflow_id": wf["id"],
                    "section_id": state.get("active_section_id"),
                    "prompt_id": prompt_id,
                    "candidate_round": candidate_round,
                    # A round number is only loop position, not semantic call
                    # identity.  Recovery may intentionally reset/rebase a
                    # budget after a prompt or contract upgrade.  Binding the
                    # key to the complete envelope preserves restart
                    # idempotency for identical input while preventing reuse of
                    # an older candidate under changed prompts or feedback.
                    "input_hash": sha256_json(envelope),
                    "model_response_protocol_version": MODEL_RESPONSE_PROTOCOL_VERSION,
                }
            )[:24]
        provider_retry = getattr(self, "_execute_prompt_with_provider_retry", None)
        if callable(provider_retry):
            result = await provider_retry(
                wf,
                state,
                prompt_id=prompt_id,
                envelope=envelope,
                call_key=requested_call_key,
            )
        else:
            result = await self.executor.execute(
                prompt_id,
                envelope,
                project_id=wf["project_id"],
                workflow_id=wf["id"],
                original_environment=state.get("original_environment"),
                call_key=requested_call_key,
            )
        if prompt_id == "P-WRITE-CONTENT" and self.diagram_enrichment is not None and result["status"] == "PASS":
            result["output"] = await self.diagram_enrichment.enrich(
                project_id=wf["project_id"],
                workflow_id=wf["id"],
                run_id=result["run_id"],
                section=section,
                output=result["output"],
                security_level=(
                    result["output"].get("source_refs", [{}])[0].get("security_level", "INTERNAL")
                    if result["output"].get("source_refs") else "INTERNAL"
                ),
            )
        result = self._apply_section_decision(
            wf,
            state,
            prompt_id,
            result,
        )
        self._append_section_run(progress, result, prompt_id=prompt_id, role=role)
        state["original_environment"] = result["route"]["environment"]
        # A freshly generated producer object is a new repair subject. Repair
        # budgets and overrides belong to the previous candidate.
        if result["status"] == "PASS":
            critic_prompt = next(
                (
                    critic
                    for critic, (producer, _phase) in self.SECTION_CRITIC_PRODUCERS.items()
                    if producer == prompt_id
                ),
                None,
            )
            if critic_prompt:
                self._supersede_repair_subject(
                    state,
                    critic_prompt=critic_prompt,
                    producer_prompt=prompt_id,
                    reason="FRESH_PRODUCER_PASS",
                )
        self._observe_quality_result(wf, state, prompt_id, result)
        self._update(wf, state=state)
        return envelope, result

    def _schedule_acceptance_regeneration(
        self,
        state: dict[str, Any],
        progress: dict[str, Any],
        critic_prompt: str,
        critic_output: dict[str, Any],
    ) -> bool:
        """Schedule a bounded new producer candidate for an acceptance run."""
        options = state.get("options") or {}
        if not bool(options.get("acceptance_run")):
            return False
        producer_phase = self.SECTION_CRITIC_PRODUCERS.get(critic_prompt)
        if not producer_phase:
            return False
        findings = [
            item
            for item in critic_output.get("findings", [])
            if isinstance(item, dict)
        ]
        if not findings or any(not item.get("repairable", False) for item in findings):
            return False

        section_id = str(state.get("active_section_id") or "")
        round_key = f"section:{section_id}:{critic_prompt}"
        rounds = state.setdefault("acceptance_regeneration_rounds", {})
        limit = max(0, min(int(options.get("acceptance_regeneration_limit", 2)), 3))
        if int(rounds.get(round_key, 0)) >= limit:
            return False
        rounds[round_key] = int(rounds.get(round_key, 0)) + 1

        producer_prompt, producer_phase_name = producer_phase
        producer_round_key = f"section:{section_id}:{producer_prompt}"
        candidate_rounds = state.setdefault("acceptance_candidate_rounds", {})
        candidate_rounds[producer_round_key] = (
            int(candidate_rounds.get(producer_round_key, 0)) + 1
        )
        self._deactivate_repair_application(state, producer_prompt)
        progress["phase"] = producer_phase_name
        progress["status"] = "RUNNING"
        progress.pop("last_error", None)
        state.pop("last_error", None)
        return True

    @staticmethod
    def _acceptance_regenerable_review_status(
        state: dict[str, Any],
        status: str,
    ) -> bool:
        if status == "REVISE":
            return True
        options = state.get("options") or {}
        return bool(
            status == "BLOCK"
            and options.get("acceptance_run")
            and options.get("allow_repairable_block_regeneration")
        )

    def _schedule_acceptance_producer_regeneration(
        self,
        state: dict[str, Any],
        progress: dict[str, Any],
        producer_prompt: str,
        producer_output: dict[str, Any],
    ) -> bool:
        """Regenerate a producer object rejected by deterministic quality checks."""
        options = state.get("options") or {}
        producer_phase = self.SECTION_PRODUCER_PHASES.get(producer_prompt)
        if not producer_phase:
            return False
        findings = [
            item
            for item in producer_output.get("findings", [])
            if isinstance(item, dict)
        ]
        if not findings or any(not item.get("repairable", False) for item in findings):
            return False

        section_id = str(state.get("active_section_id") or "")
        round_key = f"section:{section_id}:{producer_prompt}"
        rounds = state.setdefault("acceptance_regeneration_rounds", {})
        default_limit = 2
        limit = max(
            0,
            min(
                int(options.get("acceptance_regeneration_limit", default_limit)),
                3,
            ),
        )
        if int(rounds.get(round_key, 0)) >= limit:
            return False
        rounds[round_key] = int(rounds.get(round_key, 0)) + 1
        candidate_rounds = state.setdefault("acceptance_candidate_rounds", {})
        candidate_rounds[round_key] = int(candidate_rounds.get(round_key, 0)) + 1
        self._deactivate_repair_application(state, producer_prompt)
        progress["phase"] = producer_phase
        progress["status"] = "RUNNING"
        progress.pop("last_error", None)
        state.pop("last_error", None)
        return True

    async def _write_sections(self, wf: dict[str, Any], state: dict[str, Any]) -> dict[str, Any] | None:
        """Run an isolated, recoverable producer/critic/repair chain for each section.

        The chain is:
        Blueprint -> Blueprint Critic -> bounded Targeted Repair -> re-review
        -> Content -> Content Critic -> bounded Targeted Repair -> re-review
        -> Expression Polish -> Expression Critic -> bounded Targeted Repair
        -> re-review.

        Progress is persisted after every model run.  A restart re-enters the same
        phase; the Track-A deterministic call key then reuses an already committed
        response instead of duplicating model calls or artifacts.
        """
        options = state.get("options") or {}
        state["current_workflow_id"] = wf["id"]
        sections = self._target_sections(wf["project_id"], options, state)
        if bool(options.get("single_section_complete_chain")) and len(sections) != 1:
            state["last_error"] = (
                "单章节完整链要求精确选择一个章节；当前匹配 " + str(len(sections)) + " 个。"
            )
            self._update(wf, status=WorkflowStatus.BLOCKED_CONTENT.value, state=state)
            return self.get(wf["id"])

        completed = {str(item.get("section_id") or "") for item in state.get("section_results", [])}
        state.setdefault("section_results", [])
        progress_map = state.setdefault("section_progress", {})

        for section in sections:
            section_id = str(section["section_id"])
            if section_id in completed:
                continue
            state["active_section_id"] = section_id
            state["active_section_title"] = section.get("title")
            state["active_section"] = section
            progress = progress_map.setdefault(
                section_id,
                {
                    "section_id": section_id,
                    "title": section.get("title"),
                    "phase": "BLUEPRINT",
                    "status": "RUNNING",
                    "runs": [],
                },
            )
            progress.setdefault("runs", [])
            progress.setdefault("phase", "BLUEPRINT")
            progress["status"] = "RUNNING"
            self._update(wf, state=state)

            while progress["phase"] != "DONE":
                phase = str(progress["phase"])
                if phase not in self.SECTION_PHASES:
                    error = ValueError(f"未知章节阶段：{phase}")
                    return self._block_section_chain(
                        wf,
                        state,
                        section,
                        str(error),
                        configuration_error=error,
                    )
                prompt_id, next_phase = self.SECTION_PHASES[phase]
                pending_rereview = progress.get("pending_repair_rereview")
                is_pending_rereview = (
                    isinstance(pending_rereview, dict)
                    and str(pending_rereview.get("critic_prompt") or "") == prompt_id
                )
                if isinstance(pending_rereview, dict) and not is_pending_rereview:
                    error = ValueError(
                        "Persisted repair re-review checkpoint does not match "
                        f"the current section phase: expected "
                        f"{pending_rereview.get('critic_prompt')}, current {prompt_id}."
                    )
                    return self._block_section_chain(
                        wf,
                        state,
                        section,
                        str(error),
                        configuration_error=error,
                    )
                try:
                    if is_pending_rereview:
                        self._start_repair_rereview(
                            state,
                            pending_rereview,
                            critic_prompt=prompt_id,
                        )
                        self._update(wf, state=state)
                    envelope, result = await self._execute_section_prompt(
                        wf,
                        state,
                        section,
                        progress,
                        prompt_id,
                        role=(
                            "INDEPENDENT_REVIEW"
                            if is_pending_rereview
                            else ("INITIAL_REVIEW" if prompt_id.endswith("CRITIC") else "PRODUCER")
                        ),
                    )
                    if is_pending_rereview:
                        self._complete_repair_rereview(
                            state,
                            pending_rereview,
                            critic_prompt=prompt_id,
                            review_run_id=str(result.get("run_id") or "") or None,
                            status=str(result.get("status") or ""),
                        )
                        # Keep the checkpoint until the effective result and its
                        # phase transition are committed together.
                        pending_rereview["review_run_id"] = (
                            str(result.get("run_id") or "") or None
                        )
                        pending_rereview["completed_status"] = str(
                            result.get("status") or ""
                        )
                except WorkflowInputRequired:
                    raise
                except PromptExecutionError as exc:
                    return self._block_section_prompt_failure(
                        wf, state, section, exc
                    )
                except (ValueError, KeyError) as exc:
                    return self._block_section_chain(
                        wf, state, section, str(exc), configuration_error=exc
                    )

                if result["status"] == "PASS":
                    progress.pop("pending_repair_rereview", None)
                    progress["phase"] = next_phase
                    self._update(wf, state=state)
                    continue

                if result["status"] == "NEED_USER_INPUT":
                    progress.pop("pending_repair_rereview", None)
                    progress["status"] = "WAITING_GATE"
                    state["section_input_gate"] = {
                        "section_id": section_id,
                        "section_title": section.get("title"),
                        "phase": phase,
                        "next_phase": next_phase,
                        "prompt_id": prompt_id,
                        "run_id": result["run_id"],
                    }
                    self._update(wf, status="RUNNING", state=state)
                    refreshed = self.get(wf["id"])
                    self._create_gate(
                        refreshed,
                        self.pack.entry(prompt_id).get("next_human_gate")
                        or "PROJECT_GAP_RESOLUTION",
                        target_id=result["run_id"],
                        questions=result["output"].get("user_questions", []),
                    )
                    self._update(refreshed, status="WAITING_GATE", state=state)
                    return self.get(wf["id"])

                if (
                    result["status"] == "REVISE"
                    and self._schedule_acceptance_producer_regeneration(
                        state,
                        progress,
                        prompt_id,
                        result["output"],
                    )
                ):
                    progress.pop("pending_repair_rereview", None)
                    self._update(wf, status="RUNNING", state=state)
                    continue

                if (
                    prompt_id in self.SECTION_REPAIR_CRITICS
                    and result["status"] == "BLOCK"
                    and self._acceptance_regenerable_review_status(
                        state,
                        str(result["status"]),
                    )
                    and self._schedule_acceptance_regeneration(
                        state,
                        progress,
                        prompt_id,
                        result["output"],
                    )
                ):
                    progress.pop("pending_repair_rereview", None)
                    self._update(wf, status="RUNNING", state=state)
                    continue

                if result["status"] == "REVISE" and prompt_id in self.SECTION_REPAIR_CRITICS:
                    if not self._can_auto_repair(prompt_id, state):
                        if self._schedule_acceptance_regeneration(
                            state,
                            progress,
                            prompt_id,
                            result["output"],
                        ):
                            progress.pop("pending_repair_rereview", None)
                            self._update(wf, status="RUNNING", state=state)
                            continue
                        progress.pop("pending_repair_rereview", None)
                        message = (
                            f"{prompt_id} 定向修复后的独立复审返回 {result['status']}；"
                            "禁止二次自动修复或人工改正文放行。"
                            if is_pending_rereview
                            else f"{prompt_id} 在一次定向修复后仍需修改；章节修复额度已耗尽。"
                        )
                        return self._block_section_chain(
                            wf, state, section, message,
                        )
                    repaired = await self._auto_repair(wf, prompt_id, envelope, result["output"], state)
                    if not repaired:
                        progress.pop("pending_repair_rereview", None)
                        return self._block_section_chain(
                            wf, state, section, f"{prompt_id} 返回 REVISE，但没有可执行的局部修复或定向修复失败。",
                        )
                    self._append_section_run(
                        progress,
                        repaired,
                        prompt_id="P-TARGETED-REPAIR",
                        role="TARGETED_REPAIR",
                    )
                    progress["pending_repair_rereview"] = {
                        **self._repair_rereview_checkpoint(repaired),
                        "critic_prompt": prompt_id,
                    }
                    self._update(wf, state=state)
                    # Re-enter the same Critic phase.  The persisted checkpoint
                    # makes the next call an independent re-review and survives a
                    # crash before or after REREVIEW_STARTED is recorded.
                    continue

                progress.pop("pending_repair_rereview", None)
                return self._block_section_chain(
                    wf, state, section, f"{prompt_id} 返回 {result['status']}；该阶段不允许跳过或人工覆盖。",
                )

            progress["status"] = "COMPLETED"
            section_record = {
                "section_id": section_id,
                "title": section.get("title"),
                "status": "COMPLETED",
                "runs": list(progress["runs"]),
            }
            state["section_results"].append(section_record)
            state.setdefault("section_revision_findings", {}).pop(section_id, None)
            completed.add(section_id)
            self._update(wf, state=state)

        state.pop("active_section_id", None)
        state.pop("active_section_title", None)
        state.pop("active_section", None)
        state.pop("integration_repair_section_ids", None)
        state.pop("last_error", None)
        wf["current_step"] += 1
        skip_gate = bool(state.pop("skip_candidate_gate_once", False))
        self._update(wf, current_step=wf["current_step"], state=state)
        if skip_gate:
            return None
        refreshed = self.get(wf["id"])
        self._create_gate(refreshed, "CANDIDATE_REVIEW", target_id=wf["id"], questions=[])
        self._update(refreshed, status="WAITING_GATE", state=state)
        return self.get(wf["id"])

    _REPORT_SECTION_ORDER = {"BODY": 0, "ABSTRACT": 1, "CONCLUSION": 2}
    REPORT_SECTION_MAX_ATTEMPTS = 2

    @staticmethod
    def _report_section_kind(section: dict[str, Any]) -> str:
        key = str(section.get("section_key") or "").lower()
        title = str(section.get("title") or "")
        if "reference" in key or "参考资料" in title or "证据对照" in title or "文献" in title:
            return "REFERENCES"
        if "abstract" in key or "summary" in key or "摘要" in title:
            return "ABSTRACT"
        if "conclusion" in key or "结论" in title:
            return "CONCLUSION"
        return "BODY"

    async def _write_report_sections(
        self,
        wf: dict[str, Any],
        state: dict[str, Any],
    ) -> dict[str, Any] | None:
        """Write survey-report sections one at a time with persisted progress.

        Unlike the proposal section chain there is no per-section critic loop:
        body sections are written first, abstract and conclusion last, and
        reference tables are code-generated during assembly.  Progress is
        committed after every section so a provider interruption resumes from
        the failed section instead of restarting the report.
        """

        outline = self._context_result(
            wf["project_id"],
            REPORT_OUTLINE_PROMPT,
            workflow_id=wf["id"],
            exact_workflow=True,
        ) or {}
        sections = [
            section
            for section in outline.get("report_sections") or []
            if isinstance(section, dict) and str(section.get("section_key") or "").strip()
        ]
        if not sections:
            raise ValueError(
                "报告分章写作需要已确认的 P-REPORT-OUTLINE 提纲，但未找到 report_sections。"
            )
        writable = [
            (index, section)
            for index, section in enumerate(sections)
            if self._report_section_kind(section) != "REFERENCES"
        ]
        writable.sort(
            key=lambda pair: (
                self._REPORT_SECTION_ORDER[self._report_section_kind(pair[1])],
                pair[0],
            )
        )
        progress_map = state.setdefault("report_section_progress", {})
        provider_retry = getattr(self, "_execute_prompt_with_provider_retry", None)

        for _outline_index, section in writable:
            section_key = str(section["section_key"])
            progress = progress_map.get(section_key)
            if not isinstance(progress, dict):
                progress = None
            if progress and str(progress.get("status") or "") == "COMPLETED":
                continue
            attempts = int((progress or {}).get("attempts") or 0)
            if progress and str(progress.get("status") or "") == "PENDING_REVISION":
                attempts = 0
            if attempts >= self.REPORT_SECTION_MAX_ATTEMPTS:
                progress["status"] = "FAILED"
                progress["updated_at"] = utc_now()
                self._update(wf, state=state)
                continue

            active_section = {
                key: section.get(key)
                for key in (
                    "section_key",
                    "title",
                    "goal",
                    "must_answer_questions",
                    "evidence_card_ids",
                    "known_gaps",
                    "planned_exhibits",
                )
                if section.get(key) is not None
            }
            # The section-write input schema requires goal/questions/gaps; a
            # missing value would silently drop the section replacement and
            # leak the pack's placeholder section to the model.
            if not str(active_section.get("goal") or "").strip():
                active_section["goal"] = str(section.get("title") or section_key)
            active_section.setdefault("must_answer_questions", [])
            active_section.setdefault("known_gaps", [])
            state["active_report_section"] = active_section
            progress = {
                "section_key": section_key,
                "title": section.get("title"),
                "status": "RUNNING",
                "attempts": attempts,
                "updated_at": utc_now(),
            }
            progress_map[section_key] = progress

            completed = False
            while (
                not completed
                and int(progress.get("attempts") or 0) < self.REPORT_SECTION_MAX_ATTEMPTS
            ):
                progress["attempts"] = int(progress.get("attempts") or 0) + 1
                progress["status"] = "RUNNING"
                progress["updated_at"] = utc_now()
                self._update(wf, state=state)
                try:
                    envelope = self.context_builder.build(
                        REPORT_SECTION_WRITE_PROMPT,
                        wf["project_id"],
                        workflow_id=wf["id"],
                        workflow_state=state,
                    )
                    if callable(provider_retry):
                        result = await provider_retry(
                            wf,
                            state,
                            prompt_id=REPORT_SECTION_WRITE_PROMPT,
                            envelope=envelope,
                        )
                    else:
                        result = await self.executor.execute(
                            REPORT_SECTION_WRITE_PROMPT,
                            envelope,
                            project_id=wf["project_id"],
                            workflow_id=wf["id"],
                            original_environment=state.get("original_environment"),
                        )
                except (WorkflowInputRequired, ProviderRetriesExhausted):
                    # Progress is already persisted; a resume retries this section
                    # because attempts remain below the cap.
                    self._update(wf, state=state)
                    raise
                except Exception as exc:  # per-section isolation: record and continue
                    progress["last_error"] = str(exc)[:500]
                    progress["updated_at"] = utc_now()
                    self._update(wf, state=state)
                    continue
                if str(result.get("status") or "") == "PASS":
                    output_result = (result.get("output") or {}).get("result") or {}
                    body = str(output_result.get("markdown_body") or "")
                    if not body.strip():
                        progress["last_error"] = "P-REPORT-SECTION-WRITE 返回空正文"
                        progress["updated_at"] = utc_now()
                        self._update(wf, state=state)
                        continue
                    artifact_id = new_id("artifact")
                    content = {
                        "section_key": section_key,
                        "title": section.get("title"),
                        "markdown_body": body,
                        "cited_card_ids": output_result.get("cited_card_ids") or [],
                        "unresolved_questions": output_result.get("unresolved_questions") or [],
                        "run_id": str(result.get("run_id") or ""),
                        "created_at": utc_now(),
                    }
                    row = self.db.fetchone(
                        "SELECT COALESCE(MAX(version),0) AS v FROM artifacts WHERE project_id=? AND workflow_id=? AND artifact_type='REPORT_SECTION'",
                        (wf["project_id"], wf["id"]),
                    )
                    self.db.execute(
                        """INSERT INTO artifacts(id,project_id,workflow_id,artifact_type,prompt_id,version,status,security_level,context_hash,content_json,created_at)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            artifact_id,
                            wf["project_id"],
                            wf["id"],
                            "REPORT_SECTION",
                            REPORT_SECTION_WRITE_PROMPT,
                            int((row or {}).get("v") or 0) + 1,
                            "PASS",
                            self._project_level(wf["project_id"]),
                            sha256_json({k: v for k, v in content.items() if k != "created_at"}),
                            json.dumps(content, ensure_ascii=False),
                            content["created_at"],
                        ),
                    )
                    progress["status"] = "COMPLETED"
                    progress["run_id"] = content["run_id"]
                    progress["artifact_id"] = artifact_id
                    progress["summary"] = body[:200]
                    progress.pop("last_error", None)
                    progress["updated_at"] = utc_now()
                    self._update(wf, state=state)
                    completed = True
                    continue
                progress["last_error"] = (
                    f"P-REPORT-SECTION-WRITE 返回 {result.get('status') or 'ERROR'}"
                )
                progress["updated_at"] = utc_now()
                self._update(wf, state=state)

            if not completed:
                progress["status"] = "FAILED"
                progress["updated_at"] = utc_now()
                self._update(wf, state=state)

        state.pop("active_report_section", None)
        completed_keys = [
            key
            for key, progress in progress_map.items()
            if isinstance(progress, dict) and str(progress.get("status") or "") == "COMPLETED"
        ]
        if not completed_keys:
            state["last_error"] = "报告分章写作没有任何一章成功生成正文。"
            self._update(wf, status=WorkflowStatus.BLOCKED_CONTENT.value, state=state)
            return self.get(wf["id"])
        wf["current_step"] += 1
        self._update(wf, current_step=wf["current_step"], state=state)
        return None

    def _assemble_report(self, wf: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
        """Assemble the survey report Markdown and persist it.

        Assembly is deterministic code: the outline order decides chapter
        placement, reference tables are generated from the evidence catalog,
        and failed chapters remain visible as explicit placeholders instead of
        being silently dropped.  The delivery record states content and check
        status truthfully.
        """

        outline = self._context_result(
            wf["project_id"],
            REPORT_OUTLINE_PROMPT,
            workflow_id=wf["id"],
            exact_workflow=True,
        ) or {}
        sections = [
            section
            for section in outline.get("report_sections") or []
            if isinstance(section, dict) and str(section.get("section_key") or "").strip()
        ]
        if not sections:
            raise ValueError("报告整合需要已确认的 P-REPORT-OUTLINE 提纲。")
        project = self.db.fetchone(
            "SELECT name FROM projects WHERE id=?", (wf["project_id"],)
        ) or {}
        report_title = str(
            outline.get("report_title") or project.get("name") or "调研报告"
        ).strip()
        report_context = self.context_builder._wf4_report_branch_context(
            wf["project_id"], state
        ) or {}
        progress_map = (
            state.get("report_section_progress")
            if isinstance(state.get("report_section_progress"), dict)
            else {}
        )
        drafts = {
            str(draft.get("section_key")): draft
            for draft in self.context_builder._wf4_report_section_drafts(
                wf["project_id"], wf["id"], state
            )
        }
        cards = report_context.get("background_cards") or []
        sources = {
            str(source.get("source_id") or ""): source
            for source in report_context.get("source_catalog") or []
            if isinstance(source, dict) and str(source.get("source_id") or "").strip()
        }

        missing_sections: list[str] = []
        parts: list[str] = [f"# {report_title}", ""]
        parts.append("## 目录")
        parts.append("")
        for index, section in enumerate(sections, start=1):
            parts.append(f"{index}. {section.get('title') or section['section_key']}")
        parts.append("")
        for section in sections:
            section_key = str(section["section_key"])
            title = str(section.get("title") or section_key)
            parts.append(f"## {title}")
            parts.append("")
            if self._report_section_kind(section) == "REFERENCES":
                parts.extend(self._report_references_markdown(cards, sources))
            elif section_key in drafts:
                parts.append(
                    self._normalize_report_body(
                        str(drafts[section_key].get("markdown_body") or "").strip(),
                        title,
                    )
                )
            else:
                progress = progress_map.get(section_key)
                reason = (
                    str(progress.get("last_error") or "")
                    if isinstance(progress, dict)
                    else ""
                ) or "该章节未完成正文生成"
                missing_sections.append(title)
                parts.append(f"【本章未能生成：{reason}】")
            parts.append("")

        gap_lines: list[str] = []
        for gap in outline.get("overall_gaps") or []:
            text = str(gap if not isinstance(gap, dict) else gap.get("description") or "").strip()
            if text:
                gap_lines.append(f"- {text}")
        for section in sections:
            for gap in section.get("known_gaps") or []:
                text = str(gap).strip()
                if text:
                    gap_lines.append(f"- 【{section.get('title') or section['section_key']}】{text}")
        unresolved = [
            item
            for item in state.get("report_unresolved_findings") or []
            if isinstance(item, dict)
        ]
        if gap_lines or unresolved:
            parts.append("## 附录：已知缺口")
            parts.append("")
            parts.extend(gap_lines)
            for finding in unresolved:
                description = str(finding.get("description") or "").strip()
                if description:
                    parts.append(f"- 【内容检查未解决】{description}")
            parts.append("")

        critic_result = self._context_result(
            wf["project_id"],
            REPORT_CONTENT_CRITIC_PROMPT,
            workflow_id=wf["id"],
            exact_workflow=True,
        ) or {}
        if not critic_result:
            # A critic that never passed (for example REVISE kept on the record
            # after the single revision round) has no PASS artifact; read the
            # latest run output directly so the report never claims NOT_RUN
            # for a review that actually happened.
            row = self.db.fetchone(
                "SELECT output_json FROM prompt_runs WHERE project_id=? AND workflow_id=? AND prompt_id=? AND output_json IS NOT NULL ORDER BY created_at DESC LIMIT 1",
                (wf["project_id"], wf["id"], REPORT_CONTENT_CRITIC_PROMPT),
            )
            if row:
                try:
                    critic_output = json.loads(row.get("output_json") or "{}")
                except (TypeError, json.JSONDecodeError):
                    critic_output = {}
                critic_result = critic_output.get("result") or {}
                if not critic_result.get("verdict") and critic_output.get("status"):
                    critic_result = {**critic_result, "status": critic_output["status"]}
        verdict = str(
            critic_result.get("verdict")
            or critic_result.get("status")
            or "NOT_RUN"
        )
        if verdict == "REVISE" and unresolved:
            verdict = "REVISE（未解决项已如实记录）"
        checks = {
            "content_critic_verdict": verdict,
            "content_critic_unresolved_findings": len(unresolved),
            "revision_rounds_used": int(state.get("report_revision_round") or 0),
            "outbound_security_checks": "NOT_TRIGGERED（报告分支全部本地处理，无对外发送内容）",
        }
        parts.append("## 附录：检查状态")
        parts.append("")
        parts.append(f"- 全文内容检查结论：{verdict}")
        parts.append("- 外围安全/出境检查：未触发（报告分支全部本地处理，无对外发送内容）")
        if unresolved:
            parts.append(f"- 存在 {len(unresolved)} 条内容检查未解决项，已在上方缺口附录中列出。")
        if missing_sections:
            parts.append("- 以下章节未能生成正文：" + "、".join(missing_sections))
        parts.append("")

        markdown = "\n".join(parts).strip() + "\n"
        content_status = "COMPLETED" if not missing_sections else "PARTIAL"
        delivery: dict[str, Any] = {
            "content_status": content_status,
            "missing_sections": missing_sections,
            "checks": checks,
            "generated_at": utc_now(),
        }

        settings = getattr(
            getattr(getattr(self, "executor", None), "gateway", None), "settings", None
        )
        exports_dir = Path(
            getattr(settings, "exports_dir", None) or (Path("data") / "exports")
        )
        exports_dir.mkdir(parents=True, exist_ok=True)
        markdown_path = exports_dir / f"report_{wf['id']}.md"
        markdown_path.write_text(markdown, encoding="utf-8")
        delivery["markdown_path"] = str(markdown_path)

        docx_path = exports_dir / f"report_{wf['id']}.docx"
        try:
            self._write_report_docx(markdown, docx_path)
            delivery["docx_status"] = "GENERATED"
            delivery["docx_path"] = str(docx_path)
        except Exception as exc:
            delivery["docx_status"] = "FAILED"
            delivery["docx_error"] = str(exc)[:300]

        artifact_id = new_id("artifact")
        artifact_content = {
            "report_title": report_title,
            "markdown": markdown,
            "delivery": {key: value for key, value in delivery.items() if key != "docx_error"},
            "created_at": delivery["generated_at"],
        }
        row = self.db.fetchone(
            "SELECT COALESCE(MAX(version),0) AS v FROM artifacts WHERE project_id=? AND workflow_id=? AND artifact_type='REPORT_MARKDOWN'",
            (wf["project_id"], wf["id"]),
        )
        self.db.execute(
            """INSERT INTO artifacts(id,project_id,workflow_id,artifact_type,prompt_id,version,status,security_level,context_hash,content_json,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (
                artifact_id,
                wf["project_id"],
                wf["id"],
                "REPORT_MARKDOWN",
                None,
                int((row or {}).get("v") or 0) + 1,
                "PASS",
                self._project_level(wf["project_id"]),
                sha256_json({k: v for k, v in artifact_content.items() if k != "created_at"}),
                json.dumps(artifact_content, ensure_ascii=False),
                artifact_content["created_at"],
            ),
        )
        self.db.audit(
            "REPORT_ASSEMBLED",
            project_id=wf["project_id"],
            object_id=artifact_id,
            metadata={
                "workflow_id": wf["id"],
                "content_status": content_status,
                "missing_sections": missing_sections,
                "docx_status": delivery["docx_status"],
            },
        )
        state["report_markdown_artifact_id"] = artifact_id
        state["report_delivery"] = delivery
        return delivery

    @classmethod
    def _normalize_report_body(cls, body: str, title: str) -> str:
        """Normalize one section body for whole-document assembly.

        Section bodies are model output: subsection headings often reuse the
        chapter level (``##``) and figures arrive as ``[[MERMAID]]`` placeholder
        blocks.  Assembly demotes every in-body heading one level (chapter
        titles own ``##``) and rewrites mermaid placeholders into fenced
        ``mermaid`` code blocks with a bold caption.
        """

        body = cls._strip_duplicate_section_heading(body, title)
        lines = body.splitlines()
        out: list[str] = []
        index = 0
        while index < len(lines):
            line = lines[index]
            if line.strip().startswith("[[MERMAID]]"):
                caption = (
                    line.strip()[len("[[MERMAID]]"):].split("|")[0].strip() or "图示"
                )
                index += 1
                graph: list[str] = []
                while index < len(lines) and lines[index].strip():
                    graph.append(lines[index])
                    index += 1
                out.append(f"**图：{caption}**")
                out.append("")
                out.append("```mermaid")
                out.extend(graph)
                out.append("```")
                continue
            stripped = line.lstrip()
            if stripped.startswith("#"):
                hashes = len(stripped) - len(stripped.lstrip("#"))
                line = "#" * min(hashes + 1, 6) + stripped[hashes:]
            out.append(line)
            index += 1
        text = "\n".join(out)
        while "\n\n\n" in text:
            text = text.replace("\n\n\n", "\n\n")
        return text

    @staticmethod
    def _strip_duplicate_section_heading(body: str, title: str) -> str:
        """Drop a leading heading that just repeats the assembly chapter title."""

        lines = body.splitlines()
        first = next((index for index, line in enumerate(lines) if line.strip()), None)
        if first is None:
            return body
        heading = lines[first].lstrip("#").strip()
        if lines[first].lstrip().startswith("#") and heading == title.strip():
            del lines[first]
            while lines and not lines[0].strip():
                del lines[0]
            return "\n".join(lines)
        return body

    @staticmethod
    def _report_references_markdown(
        cards: list[dict[str, Any]],
        sources: dict[str, dict[str, Any]],
    ) -> list[str]:
        lines = [
            "| 证据卡 | 结论摘要 | 来源 |",
            "|---|---|---|",
        ]
        for card in cards:
            card_id = str(card.get("card_id") or "")
            claim = str(card.get("claim_text") or "").replace("|", "\\|")
            if len(claim) > 80:
                claim = claim[:80] + "…"
            source_labels = []
            for source_id in card.get("source_ids") or []:
                source = sources.get(str(source_id)) or {}
                title = str(source.get("title") or source_id)
                url = str(source.get("url") or "")
                label = f"[{source_id}] {title}"
                if url:
                    label = f"[{source_id}] [{title}]({url})"
                source_labels.append(label.replace("|", "\\|"))
            lines.append(f"| {card_id} | {claim} | {'<br>'.join(source_labels)} |")
        if len(lines) == 2:
            lines.append("| （无） | 本次调研未产生已绑定来源的证据卡 | |")
        lines.append("")
        return lines

    @staticmethod
    def _write_report_docx(markdown: str, path: Path) -> None:
        import docx

        document = docx.Document()
        in_code = False
        for line in markdown.splitlines():
            if line.strip().startswith("```"):
                if in_code:
                    document.add_paragraph("（图示源码见 Markdown 全文的 mermaid 代码块）")
                in_code = not in_code
                continue
            if in_code:
                continue
            if line.startswith("### "):
                document.add_heading(line[4:].strip(), level=3)
            elif line.startswith("## "):
                document.add_heading(line[3:].strip(), level=2)
            elif line.startswith("# "):
                document.add_heading(line[2:].strip(), level=1)
            elif line.strip():
                document.add_paragraph(line.replace("**", ""))
        document.save(str(path))

    def _three_section_mode(self, state: dict[str, Any]) -> bool:
        options = state.get("options") or {}
        return bool(
            options.get("three_section_cross_chapter")
            or options.get("integration_scope") == "THREE_SECTION_CROSS_CHAPTER"
        )

    def _resolve_three_section_contract(
        self,
        sections: list[dict[str, Any]],
        state: dict[str, Any],
    ) -> list[dict[str, Any]]:
        by_profile: dict[str, list[dict[str, Any]]] = {profile: [] for profile in THREE_SECTION_PROFILE_ORDER}
        for section in sections:
            profile = self.pack.section_profile_for(str(section.get("title") or ""))
            profile_id = str(profile.get("profile_id") or "")
            if profile_id in by_profile:
                by_profile[profile_id].append(section)
        missing = [profile for profile, values in by_profile.items() if not values]
        duplicate = [profile for profile, values in by_profile.items() if len(values) > 1]
        if missing or duplicate:
            details = []
            if missing:
                details.append("缺少章节角色：" + "、".join(missing))
            if duplicate:
                details.append("章节角色重复：" + "、".join(duplicate))
            raise ValueError(
                "三章节跨章链必须且只能包含背景、研究内容、技术路线三个唯一章节；" + "；".join(details)
            )
        resolved = [by_profile[profile][0] for profile in THREE_SECTION_PROFILE_ORDER]
        state["three_section_contract"] = {
            "contract_type": "THREE_SECTION_CROSS_CHAPTER",
            "ordered_profiles": list(THREE_SECTION_PROFILE_ORDER),
            "sections": [
                {
                    "section_id": str(section.get("section_id")),
                    "title": str(section.get("title") or ""),
                    "profile_id": profile,
                    "order": index + 1,
                }
                for index, (profile, section) in enumerate(zip(THREE_SECTION_PROFILE_ORDER, resolved))
            ],
        }
        return resolved

    def _validate_three_section_integration_envelope(
        self,
        state: dict[str, Any],
        envelope: dict[str, Any],
    ) -> None:
        if not self._three_section_mode(state):
            return
        contract = state.get("three_section_contract") or {}
        expected = [
            str(item.get("section_id")) for item in contract.get("sections") or []
            if isinstance(item, dict) and item.get("section_id")
        ]
        candidates = (envelope.get("payload") or {}).get("candidate_sections") or []
        actual = [
            str(item.get("section_id")) for item in candidates
            if isinstance(item, dict) and item.get("section_id")
        ]
        if len(expected) != 3 or len(actual) != 3 or set(actual) != set(expected):
            raise ValueError(
                "三章节跨章审查输入必须与已冻结的背景—研究内容—技术路线合同完全一致；"
                f"expected={expected}, actual={actual}"
            )

    @staticmethod
    def _section_ids_from_integration_output(
        output: dict[str, Any],
        known_section_ids: set[str],
    ) -> set[str]:
        result = output.get("result") or {}
        affected: set[str] = set()
        for key in ("redundancy_report", "document_type_drift", "page_budget_check"):
            report = result.get(key) or {}
            affected.update(str(x) for x in report.get("affected_section_ids", []) if x)
            affected.update(str(x) for x in report.get("overflow_section_ids", []) if x)
        for check in result.get("terminology_checks") or []:
            if isinstance(check, dict) and not check.get("consistent", True):
                affected.update(str(x) for x in check.get("sections", []) if x)
        evidence_strings: list[str] = []
        for check in result.get("numeric_checks") or []:
            if isinstance(check, dict) and not check.get("consistent", True):
                evidence_strings.extend(str(x) for x in check.get("occurrences", []) if x)
        for finding in output.get("findings") or []:
            if not isinstance(finding, dict):
                continue
            evidence_strings.append(str(finding.get("target_path_or_span") or ""))
            evidence_strings.extend(str(x) for x in finding.get("evidence_refs", []) if x)
        for value in evidence_strings:
            for section_id in known_section_ids:
                if section_id and section_id in value:
                    affected.add(section_id)
        return affected & known_section_ids

    def _prepare_integration_repair(self, wf: dict[str, Any], state: dict[str, Any], output: dict[str, Any]) -> str:
        """Route full-document findings to the earliest stage able to fix them.

        Argument defects require a new argument architecture; ownership and
        dependency defects require a new narrative plan; prose repetition can
        be repaired by rewriting only affected sections.  A later stage is
        never allowed to cosmetically mask an upstream structural defect.
        """
        findings = [
            item for item in output.get("findings", [])
            if isinstance(item, dict) and item.get("blocking", True)
        ]
        argument_routes = {"ARGUMENT_ARCHITECTURE_AGENT", "PROJECT_KNOWLEDGE_AGENT"}
        planning_codes = {
            "QG_DOCUMENT_DUPLICATE_INFORMATION_KEYS",
            "QG_DOCUMENT_CLAIM_OVERCONCENTRATION",
            "PAGE_BUDGET_EXCEEDED",
        }
        repairable_codes = {
            "QG_DOCUMENT_TEMPLATE_REPETITION", "DOCUMENT_TEMPLATE_REPETITION",
            "QG_DOCUMENT_DOMINATED_BY_AGENT_SYSTEM", "DOCUMENT_TYPE_DRIFT",
        }

        def effective_route(item: dict[str, Any]) -> str:
            explicit = str(item.get("suggested_route") or "").strip()
            if explicit:
                return explicit
            code = str(item.get("code") or "")
            if code in planning_codes:
                return "PLANNING_AGENT"
            if str(item.get("category") or "") == "ARGUMENT":
                return "ARGUMENT_ARCHITECTURE_AGENT"
            if code in repairable_codes:
                return "WRITING_AGENT"
            return ""

        routed_findings = [(item, effective_route(item)) for item in findings]
        argument_findings = [
            item for item, route in routed_findings if route in argument_routes
        ]
        planning_findings = [
            item for item, route in routed_findings if route == "PLANNING_AGENT"
        ]

        if argument_findings:
            rounds = int(state.get("integration_argument_rounds", 0))
            if rounds >= 1:
                state["last_error"] = "全篇审查在一次论证架构重构后仍发现上游论证缺陷，需要补充事实或由项目负责人调整中心命题。"
                self._update(wf, status=WorkflowStatus.BLOCKED_CONTENT.value, state=state)
                return "EXHAUSTED"
            state["integration_argument_rounds"] = rounds + 1
            state["argument_revision_findings"] = argument_findings
            state["planning_revision_findings"] = []
            state["section_results"] = []
            state.pop("integration_repair_section_ids", None)
            state.pop("integration_repair_findings", None)
            # Integration review is another legitimate route back to the
            # Argument Producer.  It establishes a new semantic subject just
            # like a direct ORIGINAL_PRODUCER finding, so any active local
            # Argument repair must stop masking the producer before the jump.
            self._supersede_repair_subject(
                state,
                critic_prompt="P-ARGUMENT-ARCHITECTURE-CRITIC",
                producer_prompt="P-ARGUMENT-ARCHITECTURE",
                reason="INTEGRATION_ARGUMENT_REGENERATION_SCHEDULED",
            )
            target_step = next(
                index for index, step in enumerate(self.get(wf["id"])["steps"])
                if step.get("prompt_id") == "P-ARGUMENT-ARCHITECTURE"
            )
            wf["current_step"] = target_step
            self._update(wf, status="RUNNING", current_step=target_step, state=state)
            return "SCHEDULED"

        if planning_findings:
            rounds = int(state.get("integration_planning_rounds", 0))
            if rounds >= 1:
                state["last_error"] = "全篇审查在一次章节合同重构后仍发现命题或信息归属冲突，需要人工调整论证架构。"
                self._update(wf, status=WorkflowStatus.BLOCKED_CONTENT.value, state=state)
                return "EXHAUSTED"
            state["integration_planning_rounds"] = rounds + 1
            state["planning_revision_findings"] = planning_findings
            state["section_results"] = []
            state.pop("integration_repair_section_ids", None)
            state.pop("integration_repair_findings", None)
            target_step = next(
                index for index, step in enumerate(self.get(wf["id"])["steps"])
                if step.get("prompt_id") == "P-REVISION-PLAN"
            )
            wf["current_step"] = target_step
            self._update(wf, status="RUNNING", current_step=target_step, state=state)
            return "SCHEDULED"

        contract_sections = (state.get("three_section_contract") or {}).get("sections") or []
        known_section_ids = {
            str(item.get("section_id")) for item in contract_sections
            if isinstance(item, dict) and item.get("section_id")
        }
        if not known_section_ids:
            known_section_ids = {
                str(item.get("section_id")) for item in state.get("section_results", [])
                if isinstance(item, dict) and item.get("section_id")
            }
        affected = self._section_ids_from_integration_output(output, known_section_ids)
        writing_findings = [
            item for item, route in routed_findings if route == "WRITING_AGENT"
        ]
        if writing_findings and not affected and self._three_section_mode(state):
            # A writing defect without a section locator cannot be silently accepted.
            # Regenerate the frozen three-section set rather than allowing manual edits.
            affected = set(known_section_ids)
        if not affected or not writing_findings:
            return "NOT_APPLICABLE"
        rounds = int(state.get("integration_repair_rounds", 0))
        if rounds >= 2:
            state["last_error"] = "全篇质量审查在两轮章节重写后仍未通过；需要修改论证架构或补充事实证据。"
            self._update(wf, status=WorkflowStatus.BLOCKED_CONTENT.value, state=state)
            return "EXHAUSTED"
        state["integration_repair_rounds"] = rounds + 1
        state["integration_repair_section_ids"] = sorted(affected)
        state["integration_repair_findings"] = writing_findings
        state.setdefault("cross_section_repair_history", []).append({
            "round": rounds + 1,
            "finding_codes": [str(item.get("code") or "") for item in writing_findings],
            "responsible_section_ids": sorted(affected),
            "route": "WRITING_AGENT",
        })
        state["section_results"] = [
            item for item in state.get("section_results", []) if str(item.get("section_id")) not in affected
        ]
        # S1 persists per-section phase checkpoints.  A cross-section repair must
        # invalidate only the responsible sections; otherwise a recovered workflow
        # would see phase=DONE and reuse the unchanged body without creating repair
        # evidence for the open integration finding.
        progress_map = state.setdefault("section_progress", {})
        for section_id in affected:
            progress_map.pop(str(section_id), None)
        state["skip_candidate_gate_once"] = True
        write_step = next(
            index for index, step in enumerate(self.get(wf["id"])["steps"])
            if step.get("type") == "WRITE_SECTIONS"
        )
        wf["current_step"] = write_step
        self._update(wf, status="RUNNING", current_step=write_step, state=state)
        return "SCHEDULED"

    def _target_sections(
        self,
        project_id: str,
        options: dict[str, Any],
        state: dict[str, Any] | None = None,
        *,
        workflow_id: str | None = None,
    ) -> list[dict[str, Any]]:
        source_sections = [
            section
            for section in self.context_builder.sections(project_id, "CURRENT_PROPOSAL")
            if section.get("level", 0) >= 1 and section.get("title", "").strip() not in {"", "全文"}
        ]
        by_id = {str(section.get("section_id")): section for section in source_sections if section.get("section_id")}
        by_title = {str(section.get("title")): section for section in source_sections if section.get("title")}

        # The approved narrative architecture, not the uploaded draft's raw
        # heading count, determines what belongs in the proposal.  This prevents
        # a long source outline from becoming dozens of same-type writing tasks.
        plan_workflow_id = str(
            (state or {}).get("parent_workflow_id")
            or workflow_id
            or (state or {}).get("current_workflow_id")
            or ""
        ) or None
        plan = self._context_result(
            project_id,
            "P-REVISION-PLAN",
            "revision_plan",
            workflow_id=plan_workflow_id,
            exact_workflow=bool(plan_workflow_id),
        ) or {}
        architecture = plan.get("narrative_architecture") or {}
        planned: list[dict[str, Any]] = []
        planned_ids: set[str] = set()
        project_row = self.db.fetchone("SELECT security_level FROM projects WHERE id=?", (project_id,)) or {}
        project_security_level = str(project_row.get("security_level") or "INTERNAL")
        for contract in architecture.get("section_contracts", []):
            if not isinstance(contract, dict) or contract.get("placement") == "OMIT":
                continue
            contract_title = str(contract.get("title") or "").strip()
            if contract_title in {"", "全文"}:
                continue
            section = by_id.get(str(contract.get("section_id"))) or by_title.get(contract_title)
            if section is None and contract.get("section_id") and contract.get("title"):
                # A confirmed revision plan may define a new target section even
                # when no CURRENT_PROPOSAL draft exists.  Represent that target
                # explicitly instead of falling back to a Replay sample.
                empty_text = ""
                section = {
                    "section_id": str(contract["section_id"]),
                    "section_key": str(contract.get("profile_id") or contract["section_id"]),
                    "title": str(contract["title"]),
                    "level": 1,
                    "text": empty_text,
                    "text_hash": sha256_text(empty_text),
                    "block_ids": [],
                    "contains_table": False,
                    "contains_formula": False,
                    "contains_image": False,
                    "contains_comment": False,
                    "contains_revision": False,
                    "security_level": project_security_level,
                }
            section_id = str((section or {}).get("section_id") or "")
            if section and section_id not in planned_ids:
                planned.append(section)
                planned_ids.add(section_id)
        sections = planned or source_sections
        effective_state = state if state is not None else {"options": options}
        if self._three_section_mode(effective_state):
            sections = self._resolve_three_section_contract(sections, effective_state)

        requested_ids = set(options.get("target_section_ids") or [])
        requested_titles = set(options.get("target_section_titles") or [])
        repair_ids = set((state or {}).get("integration_repair_section_ids") or [])
        if repair_ids:
            requested_ids = repair_ids
        if requested_ids or requested_titles:
            sections = [
                section
                for section in sections
                if section.get("section_id") in requested_ids or section.get("title") in requested_titles
            ]
        if sections:
            return sections
        raise WorkflowInputRequired(
            "P-WRITE-BLUEPRINT",
            gate_type=CURRENT_PROPOSAL_INPUT,
            missing_paths=["payload.source_section", "payload.revision_plan.narrative_architecture.section_contracts"],
            questions=material_input_questions(CURRENT_PROPOSAL_INPUT),
            message=(
                "没有 CURRENT_PROPOSAL 章节，也没有已确认修订计划中的明确目标章节。"
                "系统不会使用 Replay 章节代替真实写作对象。"
            ),
        )

